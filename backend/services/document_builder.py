import os
import platform
import logging
from docx import Document
from fpdf import FPDF
import markdown

logger = logging.getLogger(__name__)


def _find_system_font(font_name: str = "arial") -> tuple[str | None, str | None]:
    """Find system font files cross-platform."""
    system = platform.system()

    font_dirs = []
    if system == "Windows":
        font_dirs = [os.path.join(os.environ.get("WINDIR", "C:\\Windows"), "Fonts")]
    elif system == "Darwin":  # macOS
        font_dirs = ["/Library/Fonts", os.path.expanduser("~/Library/Fonts"), "/System/Library/Fonts"]
    else:  # Linux
        font_dirs = ["/usr/share/fonts", "/usr/local/share/fonts", os.path.expanduser("~/.fonts")]

    regular_names = [f"{font_name}.ttf", f"{font_name.lower()}.ttf", f"Arial.ttf"]
    bold_names = [f"{font_name}bd.ttf", f"{font_name}b.ttf", f"arialbd.ttf", f"Arial Bold.ttf"]

    regular_path = None
    bold_path = None

    for font_dir in font_dirs:
        if not os.path.isdir(font_dir):
            continue
        for root, _dirs, files in os.walk(font_dir):
            lower_files = {f.lower(): f for f in files}
            for name in regular_names:
                if name.lower() in lower_files and not regular_path:
                    regular_path = os.path.join(root, lower_files[name.lower()])
            for name in bold_names:
                if name.lower() in lower_files and not bold_path:
                    bold_path = os.path.join(root, lower_files[name.lower()])

    return regular_path, bold_path


class DocumentBuilder:
    @staticmethod
    def build_docx(content: str, output_path: str):
        doc = Document()
        lines = content.split('\n')

        for line in lines:
            line = line.strip()
            if not line:
                continue

            if line.startswith('# '):
                doc.add_heading(line[2:], level=1)
            elif line.startswith('## '):
                doc.add_heading(line[3:], level=2)
            elif line.startswith('### '):
                doc.add_heading(line[4:], level=3)
            elif line.startswith('- ') or line.startswith('* '):
                doc.add_paragraph(line[2:], style='List Bullet')
            elif line.startswith('> '):
                p = doc.add_paragraph(line[2:])
                p.style = 'Quote'
            else:
                # Basic bold removal for plain docx
                clean_line = line.replace('**', '')
                doc.add_paragraph(clean_line)

        doc.save(output_path)
        return output_path

    @staticmethod
    def build_pdf(content: str, output_path: str):
        pdf = FPDF()
        pdf.add_page()

        # Cross-platform font discovery for Cyrillic support
        regular_font, bold_font = _find_system_font("arial")

        try:
            if regular_font:
                pdf.add_font('Arial', '', regular_font)
                if bold_font:
                    pdf.add_font('Arial', 'B', bold_font)
                pdf.set_font('Arial', size=12)
                logger.info(f"Using system font: {regular_font}")
            else:
                pdf.set_font("helvetica", size=12)
                logger.warning("Arial font not found, using Helvetica (Cyrillic may not render correctly)")
        except Exception as e:
            logger.warning(f"Font loading error: {e}. Falling back to Helvetica.")
            pdf.set_font("helvetica", size=12)

        # Basic markdown to HTML
        html_content = markdown.markdown(content)

        try:
            # fpdf2 supports write_html
            pdf.write_html(html_content)
        except Exception as e:
            logger.warning(f"write_html failed: {e}. Falling back to plain text.")
            pdf.multi_cell(0, 10, content.replace('**', '').replace('#', ''))

        pdf.output(output_path)
        return output_path


document_builder = DocumentBuilder()
